"""Partner adapters: daily gross invoices and separate monthly net aggregates."""
from dataclasses import dataclass
from pathlib import Path
import hashlib
import re
import numpy as np
import pandas as pd
from openpyxl import load_workbook

UNITS = {"шт": "pieces", "м": "metres", "упак": "packs"}
MONTHS = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def code(value):
    return str(int(value)) if isinstance(value, (int, float)) and value == int(value) else str(value).strip()


def rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


def read_monthly(path):
    raw = rows(path)
    columns = {}
    for i, value in enumerate(raw[0]):
        s = str(value).strip().lower()
        year = re.search(r"20\d{2}", s)
        month = next((j + 1 for j, name in enumerate(MONTHS) if s.startswith(name)), None)
        if month and year:
            columns[i] = pd.Timestamp(int(year[0]), month, 1)
    if len(columns) < 12 or str(raw[0][1]).strip() != "Номенклатура.Код":
        raise ValueError(f"Unrecognized monthly schema: {path}")
    records, blanks, negatives = {}, 0, 0
    for row in raw[1:]:
        if row[1] is None:
            continue  # second header and total row
        sku = code(row[1])
        if sku in records:
            raise ValueError(f"Duplicate monthly SKU: {sku}")
        values = {}
        for i, date in columns.items():
            value = row[i]
            if value is None:
                blanks += 1
                value = 0.0  # explicit report-format assumption; never stock data
            if not isinstance(value, (float, int)) or not np.isfinite(value):
                raise ValueError(f"Invalid monthly cell: {sku}, {date}")
            negatives += value < 0
            values[date] = float(value)  # preserve signed net aggregates
        records[sku] = values
    frame = pd.DataFrame.from_dict(records, orient="index").sort_index()
    return frame, {"filename": Path(path).name, "sha256": sha256(path), "sku_count": len(frame),
                   "blank_sales_cells_assumed_zero": blanks, "negative_net_cells_preserved": int(negatives)}


@dataclass
class Panel:
    name: str
    supplier: str
    unit: str
    warehouse: str
    daily: pd.DataFrame
    monthly_2024: pd.DataFrame
    audit: dict


def load_panels(root, suppliers=("SE", "IEK")):
    panels, audits = [], {}
    for supplier in suppliers:
        folder = Path(root) / {"SE": "Systeme electric", "IEK": "IEK"}[supplier]
        sales_paths = sorted(folder.glob("Динамика*.xlsx"))
        month_paths = sorted(folder.glob("Ежемесячные продажи*.xlsx"))
        if len(sales_paths) != 1 or len(month_paths) != 1:
            raise ValueError(f"Expected exactly one transaction and one monthly file in {folder}")
        path = sales_paths[0]
        raw = rows(path)
        expected = ["Дата", "Номер", "Документ", "Код", "Номенклатура", "Ед.", "Склад", "Количество"]
        if [str(x).strip() for x in raw[0][:8]] != expected:
            raise ValueError(f"Unexpected transaction schema: {path}")
        frame = pd.DataFrame([r[:8] for r in raw[1:]], columns=["date", "id", "document", "sku", "name", "unit", "warehouse", "quantity"])
        frame["date"] = pd.to_datetime(frame["date"], format="mixed", dayfirst=True, errors="coerce").dt.normalize()
        frame["quantity"] = pd.to_numeric(frame["quantity"], errors="coerce")
        valid = frame.date.notna() & frame.sku.notna()
        invoice = frame.document.astype(str).str.startswith("Расходная накладная")
        positive = frame.quantity.gt(0) & np.isfinite(frame.quantity)
        used = valid & invoice & positive & frame.date.ge("2025-01-01")
        clean = frame.loc[used].copy()
        clean["sku"] = clean.sku.map(code)
        if clean.empty or clean.warehouse.nunique() != 1 or not set(clean.unit).issubset(UNITS):
            raise ValueError("Adapter requires one warehouse and documented units")
        # A future unit reassignment would leak identity; fail rather than guessing a conversion.
        if clean.groupby("sku").unit.nunique().max() > 1:
            raise ValueError("SKU changes unit; explicit historical conversion required")
        monthly, maudit = read_monthly(month_paths[0])
        daily_months = clean.groupby(["sku", clean.date.dt.to_period("M").dt.to_timestamp()]).quantity.sum().unstack(fill_value=0)
        common_skus = daily_months.index.intersection(monthly.index)
        complete = [m for m in daily_months.columns if m + pd.offsets.MonthEnd(0) < clean.date.max()]
        overlap_diff = daily_months.reindex(index=common_skus, columns=complete).to_numpy() - monthly.reindex(index=common_skus, columns=complete).to_numpy()
        audit = {"transactions": {"filename": path.name, "sha256": sha256(path), "input_rows": len(frame),
                 "used_positive_invoice_rows": int(used.sum()), "excluded_rows": int((~used).sum()),
                 "negative_rows": int(frame.quantity.lt(0).sum()),
                 "invalid_date_or_code_rows": int((~valid).sum()),
                 "observed_positive_quantity_by_unit": clean.groupby("unit").quantity.sum().to_dict(),
                 "data_as_of": str(clean.date.max().date()), "start_calendar": "2025-01-01"},
                 "monthly": maudit, "overlap_reconciliation": {
                     "common_skus": len(common_skus), "complete_months": len(complete),
                     "nonmatching_sku_months": int((np.abs(overlap_diff) > 1e-6).sum()),
                     "note": "Net monthly and positive invoices differ; detail takes precedence from 2025. No addition."},
                 "monthly_only_skus_no_confirmed_historical_unit": len(monthly.index.difference(clean.sku)),
                 "historical_ingestion_timestamps_available": False,
                 "client_ids_available": False, "daily_stockout_available": False}
        audits[supplier] = audit
        for unit, group in clean.groupby("unit"):
            daily = group.groupby(["sku", "date"]).quantity.sum().unstack(fill_value=0)
            daily = daily.reindex(columns=pd.date_range("2025-01-01", clean.date.max()), fill_value=0).sort_index().astype(float)
            panels.append(Panel(f"{supplier}__{UNITS[unit]}", supplier, unit, str(group.warehouse.iloc[0]), daily,
                                monthly.loc[:, monthly.columns.year == 2024].copy(), audit))
    return panels, audits
