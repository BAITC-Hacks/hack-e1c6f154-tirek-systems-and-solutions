"""Read six SE exports without manufacturing availability, prices, or daily history.

Run the aggregate-only audit with ``python -m model.lab.sources INPUT_DIRECTORY``.
The input workbooks are opened read-only and are never saved or modified.
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from model.pipeline import read_sales

MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "июн": 6,
          "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12}
PATH_PATTERNS = {
    "sales_transactions": "Динамика*.xlsx",
    "sales_monthly": "Ежемесячные продажи*.xlsx",
    "stock_monthly": "Ежемесячные остатки*.xlsx",
    "seasonality": "Сезонность*.xlsx",
    "moq": "MOQ*.xlsx",
    "current_stock_inbound": "Товар в пути*.xlsx",
}


def rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


def month(value):
    text = str(value).strip().lower()
    year = re.search(r"20\d{2}", text)
    index = next((index for name, index in MONTHS.items() if text.startswith(name)), None)
    return pd.Timestamp(int(year[0]), index, 1) if year and index else None


def number(value):
    if isinstance(value, (bool, np.bool_)):
        return None
    return float(value) if isinstance(value, (int, float, np.number)) and np.isfinite(value) else None


def _code(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    # Numeric codes have already lost meaningful leading zeros; do not guess a join.
    if not isinstance(value, str):
        raise ValueError("SKU codes must be text: numeric cells cannot safely preserve leading zeros.")
    return value.strip()


def _records(data, code_column, start, audit):
    """Drop proven identical master rows; reject conflicting SKU rows, never last-write-wins."""
    seen = {}
    audit.update(blank_rows_excluded=0, header_footer_rows_excluded=0,
                 identical_duplicate_rows_excluded=0, conflicting_duplicate_rows=0)
    for row_number, row in enumerate(data[start:], start + 1):
        if not any(value is not None for value in row):
            audit["blank_rows_excluded"] += 1
            continue
        code = _code(row[code_column])
        if code is None:
            audit["header_footer_rows_excluded"] += 1
            continue
        if code in {"Код", "Номенклатура.Код", "Код 1с", "Итого", "Всего"}:
            audit["header_footer_rows_excluded"] += 1
            continue
        if code in seen:
            if seen[code] == row:
                audit["identical_duplicate_rows_excluded"] += 1
                continue
            raise ValueError(f"Conflicting rows for SKU {code!r}; row {row_number}. No automatic overwrite.")
        seen[code] = row
        yield code, row_number, row


def read_monthly(path, code_column, blanks_are_zero):
    data = rows(path)
    if not data or code_column >= len(data[0]) or data[0][code_column] != "Номенклатура.Код":
        raise ValueError(f"Unexpected monthly SKU header in {Path(path).name}.")
    columns = {i: month(value) for i, value in enumerate(data[0]) if month(value) is not None}
    if not columns or len(set(columns.values())) != len(columns):
        raise ValueError("Monthly columns must be present and unique.")
    result = {}
    audit = {"input_rows_after_header": len(data) - 1, "negative_cells_excluded": 0,
             "blank_cells": 0, "invalid_nonblank_cells": 0,
             "blank_policy": "zero_assumption" if blanks_are_zero else "unknown"}
    for code, _, row in _records(data, code_column, 1, audit):
        values = {}
        for index, date in columns.items():
            raw = row[index]
            value = number(raw)
            if raw is None:
                audit["blank_cells"] += 1
                value = 0.0 if blanks_are_zero else None
            elif value is None:
                audit["invalid_nonblank_cells"] += 1
            elif value < 0:
                audit["negative_cells_excluded"] += 1
                value = None
            values[date] = value
        result[code] = values
    table = pd.DataFrame.from_dict(result, orient="index", dtype=float).sort_index().sort_index(axis=1)
    audit.update(sku_count=len(table), month_count=len(columns),
                 first_month=str(min(columns.values()).date()), last_month=str(max(columns.values()).date()),
                 unknown_cells=int(table.isna().sum().sum()),
                 availability_policy="A monthly value is available only after its month ends; publication delay unknown.")
    return table, audit


def _read_moq(path):
    data, result, audit = rows(path), {}, {"nonpositive_or_missing_multiple": 0}
    if data[0][2:5] != ("Номенклатура.Код", "Артикул", "Кратность"):
        raise ValueError("Unexpected MOQ headers.")
    for code, _, row in _records(data, 2, 1, audit):
        value = number(row[4])
        result[code] = value if value is not None and value > 0 else None
        audit["nonpositive_or_missing_multiple"] += int(result[code] is None)
    audit.update(sku_count=len(result), interpretation="Кратность means order multiple; minimum order quantity is not supplied.",
                 historical_effective_dates_available=False)
    return result, audit


def _read_current(path, multiples, units):
    data, result, audit = rows(path), {}, {}
    expected = {2: "Код 1с", 4: "Категория 2026", 49: "Остаток", 50: "Зарезервировано", 51: "Свободный остаток"}
    if len(data) < 2 or any(data[1][i] != value for i, value in expected.items()):
        raise ValueError("Unexpected inventory snapshot headers.")
    match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", Path(path).name)
    if not match:
        raise ValueError("Inventory snapshot date is missing from its filename.")
    snapshot = pd.Timestamp(int(match[3]), int(match[2]), int(match[1]))
    eta_match = re.fullmatch(r"СЭ в пути (\d{2})\.(\d{2})", str(data[1][54]).strip())
    if not eta_match:
        raise ValueError("Inbound ETA is missing or ambiguous in its column header.")
    # The header lacks a year. Do not silently roll a past day/month into next year.
    eta = pd.Timestamp(snapshot.year, int(eta_match[2]), int(eta_match[1]))
    eta_verified = eta >= snapshot
    audit.update(quantity_missing_cells=0, negative_quantity_cells=0, free_stock_identity_checked=0,
                 free_stock_identity_conflicts=0, component_sum_equals_on_hand=0)
    for code, row_number, row in _records(data, 2, 2, audit):
        quantities = [number(row[i]) for i in (49, 50, 51, 54)]
        audit["quantity_missing_cells"] += sum(value is None for value in quantities)
        audit["negative_quantity_cells"] += sum(value is not None and value < 0 for value in quantities)
        on_hand, reserved, free, inbound = quantities
        consistent = None
        if all(value is not None for value in (on_hand, reserved, free)):
            consistent = bool(np.isclose(on_hand - reserved, free, rtol=0, atol=1e-8))
            audit["free_stock_identity_checked"] += 1
            audit["free_stock_identity_conflicts"] += int(not consistent)
        parts = [number(row[i]) for i in range(45, 49)]
        if on_hand is not None and all(value is not None for value in parts):
            audit["component_sum_equals_on_hand"] += int(np.isclose(sum(parts), on_hand, rtol=0, atol=1e-8))
        result[code] = {
            "sku": code, "supplier_id": "systeme-electric", "unit": units.get(code),
            "category": str(row[4]).strip() if row[4] is not None else None,
            "category_semantics_verified": False,
            "on_hand": on_hand, "reserved": reserved, "free_stock": free,
            "free_stock_consistent": consistent,
            "inbound_quantity": inbound, "inbound_eta": str(eta.date()) if eta_verified else None,
            "inbound_eta_year_assumed_from_snapshot": True,
            "inventory_as_of": str(snapshot.date()), "warehouse_scope_verified": False,
            "order_multiple": multiples.get(code), "min_order_qty": None, "lead_time_days": None,
            "reported_cost_unverified": number(row[5]), "unit_cost": None, "unit_sale_price": None,
            "reported_historical_growth": number(row[43]), "forward_growth_forecast": None,
            "source": f"{Path(path).name}:TDSheet:{row_number}",
        }
    audit.update(sku_count=len(result), as_of=str(snapshot.date()), inbound_eta=str(eta.date()) if eta_verified else None,
                 inbound_eta_year_policy="Same year as dated snapshot; unverified assumption, never automatic year rollover.",
                 category_counts=dict(Counter(record["category"] for record in result.values())),
                 unit_counts=dict(Counter(record["unit"] or "unknown" for record in result.values())),
                 missing_multiple=sum(record["order_multiple"] is None for record in result.values()),
                 negative_quantities_policy="Retained and flagged; require review before procurement.",
                 warehouse_components_policy="Do not sum component columns or add them to Остаток: their accounting scope is unverified.")
    return result, audit


def _read_turnover(path):
    result = {}
    for row in rows(path):
        if len(row) < 13 or not isinstance(row[0], int) or not 2000 <= row[0] <= 2099:
            continue
        for index in range(1, 13):
            date, value = pd.Timestamp(row[0], index, 1), number(row[index])
            if date in result:
                raise ValueError("Duplicate raw seasonality year rows.")
            if value is not None and value >= 0:
                result[date] = value
    if not result:
        raise ValueError("No raw annual seasonality rows found.")
    return result


def _transaction_quality(path):
    data = rows(path)[1:]
    counts = Counter(data)
    valid = [row for row in data if row[3] is not None]
    return {"exact_duplicate_row_excess": sum(count - 1 for count in counts.values() if count > 1),
            "repeated_document_sku_groups": sum(count > 1 for count in Counter((row[1], row[3]) for row in valid).values()),
            "numeric_sku_rows": sum(not isinstance(row[3], str) for row in valid),
            "leading_zero_sku_count": len({row[3] for row in valid if isinstance(row[3], str) and row[3].startswith("0")}),
            "duplicate_policy": "No unique line ID exists. Same document/SKU can have several real lines; these are not deduplicated.",
            "duplicate_exact_policy": "Exact repeated transactions are flagged for review, not silently removed without a unique line ID."}


def _overlap_audit(daily, monthly):
    common = daily.index.intersection(monthly.index)
    totals = daily.T.groupby(daily.columns.to_period("M")).sum().T
    compared = matched = 0
    sum_daily = sum_monthly = absolute = 0.0
    last_complete = daily.columns.max().to_period("M") - 1
    for period in totals.columns:
        stamp = period.start_time
        if period > last_complete or stamp not in monthly.columns:
            continue
        observed = totals[period].reindex(common).to_numpy(dtype=float)
        reported = monthly[stamp].reindex(common).to_numpy(dtype=float)
        finite = np.isfinite(reported)
        x, y = observed[finite], reported[finite]
        compared += len(x)
        matched += int(np.isclose(x, y, rtol=0, atol=1e-8).sum())
        sum_daily += float(x.sum())
        sum_monthly += float(y.sum())
        absolute += float(np.abs(x - y).sum())
    return {"common_skus": len(common), "compared_complete_sku_months": compared,
            "matching_sku_months": matched, "differing_sku_months": compared - matched,
            "sum_positive_daily_invoices": sum_daily, "sum_valid_monthly_report": sum_monthly,
            "sum_absolute_difference": absolute,
            "semantic_equivalence_verified": compared > 0 and matched == compared,
            "policy": "Monthly net reports and positive invoice targets have different semantics/scope. Never sum overlapping histories or replace daily labels with monthly totals."}


def load_se(directory):
    directory = Path(directory)
    paths = {}
    for role, pattern in PATH_PATTERNS.items():
        matches = sorted(directory.glob(pattern))
        if len(matches) != 1:
            raise ValueError(f"Expected one file for {role}, found {len(matches)}.")
        paths[role] = matches[0]
    daily, transactions_audit = read_sales(paths["sales_transactions"])
    transactions_audit.update(_transaction_quality(paths["sales_transactions"]))
    if transactions_audit["numeric_sku_rows"]:
        raise ValueError("Numeric transaction SKUs need explicit leading-zero mapping before joining exports.")
    monthly, monthly_audit = read_monthly(paths["sales_monthly"], 1, True)
    stock, stock_audit = read_monthly(paths["stock_monthly"], 2, False)
    units = {}
    for row in rows(paths["stock_monthly"])[1:]:
        code = _code(row[2])
        if code in stock.index:
            unit = str(row[3]).strip() if row[3] is not None else None
            if code in units and units[code] != unit:
                raise ValueError("Conflicting stock units for the same SKU.")
            units[code] = unit
    for code in daily.index:
        if code in units and units[code] not in (None, transactions_audit["unit"]):
            raise ValueError("Transaction and monthly stock units conflict.")
        units[code] = transactions_audit["unit"]
    moq, moq_audit = _read_moq(paths["moq"])
    current, current_audit = _read_current(paths["current_stock_inbound"], moq, units)
    turnover = _read_turnover(paths["seasonality"])
    source_audit = [{"role": role, "filename": path.name,
                     "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for role, path in paths.items()]
    daily_skus = set(daily.index)
    coverage = {}
    for role, mapping in (("monthly_sales", monthly.index), ("monthly_stock", stock.index),
                          ("order_multiple", moq), ("current_profile", current)):
        other = set(mapping)
        coverage[role] = {"available_skus": len(other), "daily_skus_matched": len(daily_skus & other),
                          "daily_skus_missing": len(daily_skus - other), "additional_skus": len(other - daily_skus)}
    embedded = {_code(row[1]): number(row[3]) for row in rows(paths["sales_monthly"])[1:] if row[1] is not None}
    moq_audit["embedded_sales_multiple_conflicts"] = sum(
        embedded[code] != value for code, value in moq.items() if code in embedded)
    moq_audit["conflict_policy"] = "Use the separate Кратность export; embedded sales multiples are not silently substituted. Effective date unverified."
    audit = {"transactions": transactions_audit, "monthly_sales": monthly_audit, "monthly_stock": stock_audit,
             "sources": source_audit, "current_profiles": len(current), "moq_profiles": len(moq),
             "current_snapshot": current_audit, "order_multiples": moq_audit, "coverage": coverage,
             "daily_monthly_reconciliation": _overlap_audit(daily, monthly),
             "seasonality": {"raw_months": len(turnover), "unit": "not specified in the workbook; likely monetary, unverified",
                             "feature": "Dimensionless index from a complete prior year only; precomputed cross-year coefficients ignored."},
             "capabilities": {
                 "observed_positive_sales_forecast": True, "monthly_2024_seasonality_feature": True,
                 "current_stock_and_dated_inbound": True, "order_multiples": True,
                 "stockout_lost_demand_identifiable": False, "client_level_oneoff_identifiable": False,
                 "supplier_lead_time_available": False, "verified_margin_and_holding_cost_available": False,
                 "material_obligations_available": False, "live_source_sync_available": False},
             "notes": [
                 "Monthly observations remain monthly; no fictional daily sales are created or added to daily transactions.",
                 "Monthly stock is a lagged proxy, not exact daily availability or proof of stockout.",
                 "Monthly publication/revision timestamps are unavailable. Point-in-time tests assume publication at month end; strict historical vintages cannot be verified.",
                 "Current category, stock, inbound, multiples and growth do not enter historical demand features.",
                 "Reported historical growth is not an independent forward growth forecast.",
                 "Category labels have no supplied meanings or service-level policy.",
                 "Current warehouse scope and the meaning/currency of СС реал require confirmation before cost optimization.",
                 "Client ID, selling price, daily stockout intervals, supplier lead times and material obligations are absent.",
                 "All six sources are read and audited; unavailable semantics are not invented to claim all requirements pass."]}
    return {"daily": daily, "monthly": monthly, "monthly_stock": stock, "moq": moq,
            "current": current, "units": units, "turnover": turnover, "audit": audit}


def current_profiles_at(sources, as_of):
    """Expose snapshot data only from its timestamp onwards; freshness remains caller policy."""
    cutoff = pd.Timestamp(as_of).normalize()
    return {code: deepcopy(record) for code, record in sources["current"].items()
            if pd.Timestamp(record["inventory_as_of"]) <= cutoff}


def monthly_prior_window(table, skus, origin, horizon=28):
    """Prior-year monthly prorating is a feature, never a reconstructed daily target."""
    if not isinstance(horizon, int) or horizon < 1:
        raise ValueError("Horizon must be a positive integer.")
    origin = pd.Timestamp(origin).normalize()
    if pd.isna(origin):
        raise ValueError("Origin is required.")
    result = np.zeros(len(skus))
    covered = np.zeros(len(skus))
    for future in pd.date_range(origin + pd.Timedelta(days=1), periods=horizon):
        old = future - pd.DateOffset(years=1)
        period = old.to_period("M")
        if period.end_time.normalize() > origin or period.start_time not in table.columns:
            continue
        vector = table[period.start_time].reindex(skus).to_numpy(dtype=float)
        good = np.isfinite(vector) & (vector >= 0)
        result[good] += vector[good] / period.days_in_month
        covered[good] += 1
    result[covered < horizon] = np.nan
    return result


def enrich_features(frame, sources):
    """Enrich each origin independently, without current snapshots or future-month values."""
    frame = frame.copy()
    required = {"sku", "origin", "sum84", "sum7", "sum14", "sum28", "sum56", "robust_sum28", "prior_year28"}
    if not required.issubset(frame.columns):
        raise ValueError(f"Missing base features: {sorted(required - set(frame.columns))}")
    origins = pd.to_datetime(frame["origin"], errors="raise").dt.normalize()
    if origins.isna().any():
        raise ValueError("Every feature row must have an origin.")
    for name in ("monthly_year28", "last_month_stock", "supplier_turnover_season"):
        frame[name] = np.nan
    for origin in origins.unique():
        origin = pd.Timestamp(origin)
        positions = np.flatnonzero((origins == origin).to_numpy())
        skus = frame.iloc[positions]["sku"]
        frame.iloc[positions, frame.columns.get_loc("monthly_year28")] = monthly_prior_window(sources["monthly"], skus, origin)
        previous_month = origin.to_period("M").start_time - pd.DateOffset(months=1)
        if previous_month in sources["monthly_stock"].columns:
            values = sources["monthly_stock"][previous_month].reindex(skus).to_numpy(dtype=float)
            values[values < 0] = np.nan
            frame.iloc[positions, frame.columns.get_loc("last_month_stock")] = values
        amounts = np.array([sources["turnover"].get(pd.Timestamp(origin.year - 1, m, 1), np.nan) for m in range(1, 13)])
        if np.isfinite(amounts).all() and (amounts >= 0).all() and amounts.mean() > 0:
            target_month = (origin + pd.Timedelta(days=14)).month
            frame.iloc[positions, frame.columns.get_loc("supplier_turnover_season")] = amounts[target_month - 1] / amounts.mean()
    denominator = np.maximum(frame["sum84"] / 3, 10)
    for field in ("sum7", "sum14", "sum28", "sum56", "robust_sum28", "prior_year28", "monthly_year28"):
        frame[field + "_relative"] = frame[field] / denominator
    return frame


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", type=Path)
    parser.add_argument("--audit-output", type=Path)
    args = parser.parse_args()
    payload = json.dumps(load_se(args.directory)["audit"], ensure_ascii=False, indent=2, allow_nan=False) + "\n"
    if args.audit_output:
        args.audit_output.write_text(payload, encoding="utf-8")
    else:
        print(payload, end="")


if __name__ == "__main__":
    main()
