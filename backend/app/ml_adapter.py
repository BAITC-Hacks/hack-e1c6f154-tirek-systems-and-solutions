"""Built-in bridge from uploaded SE workbooks to forecast v2 and decision core.

The adapter keeps partner workbooks in DATA_DIR, records their provenance and
returns the existing public API contract. It does not call an LLM or send orders.
"""
from copy import deepcopy
from dataclasses import replace
from datetime import date, timedelta
from decimal import Decimal
from hashlib import sha256 as hashlib_sha256
import json
import os
from pathlib import Path
import shutil

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from model.decision.core import (
    EconomicProfile, GrowthAdjustment, Inbound, InventorySnapshot,
    MaterialRequirement, Recommendation, RecommendationInput, ServicePolicy,
    SupplierConstraint, recommend,
)
from model.forecast_v2.data import load_panels, sha256
from model.forecast_v2.features import features_at
from model.forecast_v2.methods import MODEL_SPECS, load_model
from model.forecast_v2.run import forecast_methods, verify_prediction_implementation
from model.lab.sources import PATH_PATTERNS, load_se
from model.regular_forecast.predictor import VERSION as REGULAR_VERSION, forecast_with_audit

from .contracts import DomainError, ROOT
from .demo import now, summary


MODEL_DIR = ROOT / "model" / "forecast_v2" / "artifacts"
ROLE_NAMES = {
    "sales_transactions", "sales_monthly", "stock_monthly", "seasonality",
    "moq", "current_stock_inbound",
}


def _json_default(value):
    if isinstance(value, (np.integer, np.floating)):
        return value.item()
    raise TypeError(f"Cannot encode {type(value).__name__}")


def _write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2,
                                    allow_nan=False, default=_json_default) + "\n", encoding="utf-8")


def _issue(code, severity, message, skus=(), reference=None):
    return {"code": code, "severity": severity, "message": message,
            "affected_skus": list(skus), "source_reference": reference}


def _prepare_model_inputs(directory, report):
    """Create local aliases with original safe names; source bytes stay unchanged."""
    target = Path(directory) / "model_inputs" / "Systeme electric"
    target.mkdir(parents=True, exist_ok=True)
    by_role = {source["role"]: source for source in report["sources"]}
    if set(by_role) != ROLE_NAMES:
        raise ValueError("All six Systeme Electric source roles are required for normalization")
    for role, source in by_role.items():
        original = Path(source["filename"]).name
        if not original or original in (".", ".."):
            raise ValueError("Unsafe original source filename")
        if not Path(original).match(PATH_PATTERNS[role]):
            raise ValueError(
                f"Filename for {role} must preserve the partner export pattern "
                f"{PATH_PATTERNS[role]!r}; the inventory filename must contain its snapshot date"
            )
        source_path, alias = Path(directory) / f"{role}.xlsx", target / original
        if sha256(source_path) != source["sha256"]:
            raise ValueError(f"Stored source hash differs for {role}")
        if alias.exists():
            if sha256(alias) != source["sha256"]:
                raise ValueError(f"Existing model input alias differs for {role}")
            continue
        try:
            os.link(source_path, alias)
        except OSError:
            shutil.copy2(source_path, alias)
    return target.parent


def normalize_dataset(directory: Path, report: dict, context: dict | None) -> dict:
    """Normalize and audit all six uploaded workbooks for the platform import job."""
    if report['supplier_ids'] == ['iek']:
        from .iek_adapter import normalize_dataset as normalize_iek
        return normalize_iek(directory, report, context)
    directory = Path(directory)
    roles = {source["role"] for source in report["sources"]}
    if roles != ROLE_NAMES:
        report["issues"] = [issue for issue in report["issues"] if issue["code"] != "NORMALIZATION_REQUIRED"]
        report["issues"].append(_issue(
            "NORMALIZATION_BLOCKED", "error",
            "Для модели нужны все шесть ролей файлов; набор сохранён, расчёт не разрешён.",
            reference="backend.app.ml_adapter"))
        report["calculation_allowed"] = False
        return report
    try:
        model_root = _prepare_model_inputs(directory, report)
        sources = load_se(model_root / "Systeme electric")
    except (ValueError, OSError) as exc:
        raise DomainError("INVALID_FILE", f"Нормализация не завершена: {exc}", 422) from exc

    audit = sources["audit"]
    normalized = directory / "normalized"
    normalized.mkdir(exist_ok=True)
    _write_json(normalized / "current.json", sources["current"])
    _write_json(normalized / "audit.json", audit)
    _write_json(normalized / "context.json", context or {})
    _write_json(normalized / "metadata.json", {
        "model_input_root": str(model_root.relative_to(directory)),
        "data_as_of": audit["current_snapshot"]["as_of"],
        "forecast_sales_as_of": audit["transactions"]["data_as_of"],
        "synthetic_context": any(
            row.get("source_kind") == "synthetic"
            for key in ("stockout_intervals", "price_observations", "material_requirements")
            for row in (context or {}).get(key, [])
        ),
        "normalizer": "backend.app.ml_adapter:v1",
    })

    used = {
        "sales_transactions": audit["transactions"]["used_rows"],
        "sales_monthly": audit["monthly_sales"]["sku_count"],
        "stock_monthly": audit["monthly_stock"]["sku_count"],
        # One source row per year with retained raw observations. raw_months
        # counts cells and must not be presented as a number of workbook rows.
        "seasonality": len({pd.Timestamp(stamp).year for stamp in sources["turnover"]}),
        "moq": audit["moq_profiles"],
        "current_stock_inbound": audit["current_profiles"],
    }
    dated = {
        "sales_transactions": audit["transactions"]["data_as_of"],
        "current_stock_inbound": audit["current_snapshot"]["as_of"],
    }
    for source in report["sources"]:
        source["rows_used"] = int(used[source["role"]])
        source["data_as_of"] = dated.get(source["role"])

    missing_profiles = audit["coverage"]["current_profile"]["daily_skus_missing"]
    report.update(
        dataset_version=hashlib_sha256((report["dataset_version"] + "|normalized-v1").encode()).hexdigest()[:12],
        data_as_of=audit["current_snapshot"]["as_of"],
        sku_count=audit["transactions"]["sku_count"],
        calculation_allowed=True,
    )
    report["issues"] = [issue for issue in report["issues"]
                        if issue["code"] not in {"NORMALIZATION_REQUIRED", "SOURCE_DATE_UNKNOWN"}]
    context = context or {}
    synthetic_context = any(
        row.get("source_kind") == "synthetic"
        for key in ("stockout_intervals", "price_observations", "material_requirements")
        for row in context.get(key, [])
    )
    reconciliation = audit["daily_monthly_reconciliation"]
    unknown_units = audit["current_snapshot"]["unit_counts"].get("unknown", 0)
    missing_multiples = audit["current_snapshot"]["missing_multiple"]
    report["issues"].extend([
        _issue("CURRENT_STOCK_PARTIAL", "warning",
               f"У {missing_profiles} продававшихся SKU нет текущего профиля; для них количество останется неизвестным.",
               reference="normalized/audit.json"),
        _issue("CURRENT_PROFILE_WITHOUT_HISTORY", "warning",
               f"{audit['coverage']['current_profile']['additional_skus']} SKU текущего снимка не имеют дневной истории "
               "в подтверждённой единице: они будут показаны с forecast=null и needs_data.",
               reference="normalized/audit.json"),
        _issue("CONTEXT_STOCKOUT_READY" if context.get("stockout_intervals") else "NO_DAILY_STOCKOUT", "info" if context.get("stockout_intervals") else "warning",
               "Переданные интервалы stockout будут применены отдельной exposure-aware моделью; результат останется needs_review."
               if context.get("stockout_intervals") else
               "Точных интервалов отсутствия товара нет; скрытый спрос на реальных данных не восстановлен.",
               reference="normalized/audit.json"),
        _issue("CONTEXT_CLIENT_LABELS_READY" if context.get("client_labels") else "NO_CLIENT_LABELS", "info" if context.get("client_labels") else "warning",
               "Переданные клиентские метки будут связаны с номером документа или ID «номер:строка» и применены к регулярному спросу."
               if context.get("client_labels") else
               "В исходной выгрузке нет обезличенных клиентов; клиентские разовые заказы нельзя подтвердить на реальных данных.",
               reference="normalized/audit.json"),
        _issue("LEAD_TIME_FROM_REQUEST", "info",
               "Срок новой поставки берётся из параметров расчёта, поскольку справочник lead time не предоставлен.",
               reference="CalculationRequest.lead_time_days"),
        _issue("DAILY_MONTHLY_MISMATCH", "warning",
               f"Дневная и месячная выгрузки расходятся в {reconciliation['differing_sku_months']} из "
               f"{reconciliation['compared_complete_sku_months']} сопоставимых SKU-месяцев; они не суммируются.",
               reference="normalized/audit.json"),
        _issue("NEGATIVE_SALES_EXCLUDED", "warning",
               f"Нормализатор исключил {audit['transactions']['negative_rows']} отрицательных строк продаж и "
               f"{audit['monthly_sales']['negative_cells_excluded']} отрицательных месячных ячеек из своего представления. "
               "Отдельный frozen forecast v2 сохраняет отрицательные месячные net-значения 2024 года в источнике "
               "и ограничивает их вклад нулём только в агрегированном прогнозном признаке; его целевые продажи не меняются.",
               reference="normalized/audit.json"),
        _issue("UNKNOWN_UNIT", "warning", f"У {unknown_units} текущих SKU не подтверждена единица.",
               reference="normalized/audit.json"),
        _issue("UNKNOWN_ORDER_MULTIPLE", "warning", f"У {missing_multiples} текущих SKU неизвестна кратность.",
               reference="normalized/audit.json"),
        _issue("UNKNOWN_MIN_ORDER_QTY", "warning", "MOQ не предоставлен ни для одного SKU; утверждение партии блокируется.",
               reference="normalized/audit.json"),
        _issue("UNVERIFIED_WAREHOUSE_SCOPE", "warning",
               "Область складского снимка не подтверждена; количество остаётся предварительным.",
               reference="normalized/audit.json"),
        _issue("UNVERIFIED_COST", "warning",
               "Поле себестоимости исходного снимка не имеет подтверждённого смысла и не используется.",
               reference="normalized/audit.json"),
        _issue("SOURCE_EFFECTIVE_DATE_UNKNOWN", "warning",
               "Для месячных источников, сезонности и кратности нет дат публикации/вступления в силу.",
               reference="normalized/audit.json"),
        _issue("SYNTHETIC_CONTEXT_BLOCKS_OPERATIONAL", "warning",
               "Синтетический дополнительный контекст разрешён только в сценарном режиме.",
               reference="normalized/context.json") if synthetic_context else
        _issue("CATEGORY_POLICY_REQUIRED", "info",
               "Для SKU без экономического профиля нужна явная политика категории.",
               reference="CalculationRequest.category_policies"),
    ])
    if synthetic_context:
        report["issues"].append(_issue(
            "CATEGORY_POLICY_REQUIRED", "info",
            "Для SKU без экономического профиля нужна явная политика категории.",
            reference="CalculationRequest.category_policies"))
    return report


def _load_forecast(panel, origin):
    selection = json.loads((MODEL_DIR / "selection.json").read_text(encoding="utf-8"))
    lock = json.loads((MODEL_DIR / "selection.lock.json").read_text(encoding="utf-8"))
    manifest = json.loads((MODEL_DIR / "manifest.json").read_text(encoding="utf-8"))
    if sha256(MODEL_DIR / "selection.json") != lock["selection_sha256"]:
        raise ValueError("Forecast selection lock mismatch")
    verify_prediction_implementation(selection)
    if manifest["selection_sha256"] != lock["selection_sha256"] or manifest["model_specs"] != MODEL_SPECS:
        raise ValueError("Forecast manifest is incompatible with frozen selection")
    entry = manifest["panels"][panel.name]
    if pd.Timestamp(origin) < pd.Timestamp(entry["trained_as_of"]):
        raise ValueError("Final weights cannot be used before their training cutoff")
    current = features_at(panel, pd.Timestamp(origin))
    weights = {}
    for name, meta in entry["models"].items():
        path = MODEL_DIR / meta["file"]
        if sha256(path) != meta["sha256"]:
            raise ValueError(f"Forecast weight hash mismatch: {meta['file']}")
        weights[name] = load_model(name, path)
    predictions, _ = forecast_methods(panel, current, [entry["selected"], entry["best_ml"], "v1_baseline"],
                                      pd.DataFrame(), pd.DataFrame(), 1, 1, weights)
    return current, predictions[entry["selected"]], entry, lock["selection_sha256"]


def _daily_path(panel, sku, origin, total):
    """Allocate a 28-day total by observed weekday shape; this is a model trajectory."""
    history = panel.daily.loc[sku, :pd.Timestamp(origin)].iloc[-84:]
    weekday = history.groupby(history.index.weekday).mean().reindex(range(7), fill_value=0).to_numpy(float)
    future = pd.date_range(pd.Timestamp(origin) + pd.Timedelta(days=1), periods=28)
    weights = np.asarray([weekday[d.weekday()] for d in future], dtype=float)
    if not np.isfinite(weights).all() or weights.sum() <= 0:
        weights = np.ones(28)
    values = weights / weights.sum() * float(total)
    return values.tolist()


def _context_maps(context, as_of, warehouse_id):
    context = context or {}
    prices = {}
    for row in context.get("price_observations", []):
        if row["effective_date"] <= as_of and (row["sku"] not in prices or row["effective_date"] > prices[row["sku"]]["effective_date"]):
            prices[row["sku"]] = row
    materials = {}
    for row in context.get("material_requirements", []):
        if row["warehouse_id"] == warehouse_id:
            materials.setdefault(row["sku"], []).append(row)
    return prices, materials


def _transaction_events(path, labels, origin):
    """Map pseudonymous labels to positive invoice rows without exposing clients.

    A label may reference the exported document number (applies to all its rows)
    or the stable ``document:excel-row`` ID written to the audit output.
    """
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        expected = ("Дата", "Номер", "Документ", "Код", "Номенклатура", "Ед.", "Склад", "Количество")
        if tuple(str(value).strip() for value in header[:8]) != expected:
            raise ValueError("Unexpected transaction schema for client label mapping")
        events, matched = {}, set()
        for row_number, row in enumerate(rows, 2):
            if len(row) < 8 or row[1] is None or row[3] is None:
                continue
            stamp = pd.to_datetime(row[0], format="mixed", dayfirst=True, errors="coerce")
            quantity = row[7]
            if (pd.isna(stamp) or stamp.normalize() > pd.Timestamp(origin)
                    or not str(row[2]).startswith("Расходная накладная")
                    or isinstance(quantity, bool) or not isinstance(quantity, (int, float))
                    or not np.isfinite(quantity) or quantity <= 0 or str(row[5]).strip() != "шт"):
                continue
            document = str(row[1]).strip()
            event_id = f"{document}:{row_number}"
            label_key = event_id if event_id in labels else document if document in labels else None
            if label_key is None:
                continue
            matched.add(label_key)
            sku = str(row[3]).strip()
            events.setdefault(sku, []).append({
                "event_id": event_id, "date": str(stamp.date()),
                "client_id": labels[label_key], "quantity": float(quantity),
            })
        return events, matched
    finally:
        workbook.close()


def _context_forecasts(panel, origin, context, transaction_path):
    """Return reviewed daily forecasts only for SKUs affected by supplied context."""
    context = context or {}
    intervals = [row for row in context.get("stockout_intervals", [])
                 if row["warehouse_id"] in {"almaty", "Алматы"}]
    label_rows = context.get("client_labels", [])
    labels = {}
    for row in label_rows:
        key = row["source_event_id"]
        if key in labels and labels[key] != row["pseudonymous_client_id"]:
            raise DomainError("INVALID_PARAMETERS", f"Для события {key} переданы разные клиентские метки.")
        labels[key] = row["pseudonymous_client_id"]
    if labels and transaction_path is None:
        raise DomainError("MISSING_CRITICAL_DATA", "Нет источника строк продаж для клиентских меток.")
    events, matched = _transaction_events(transaction_path, labels, origin) if labels else ({}, set())
    unmatched = sorted(set(labels) - matched)
    if unmatched:
        sample = ", ".join(unmatched[:3])
        raise DomainError("MISSING_CRITICAL_DATA", f"Клиентские метки не связаны с продажами до даты расчёта: {sample}.")
    interval_skus = {row["sku"] for row in intervals}
    affected = sorted((interval_skus | set(events)) & set(panel.daily.index.astype(str)))
    daily, audits = {}, {}
    for sku in affected:
        series = panel.daily.loc[sku, :pd.Timestamp(origin)]
        availability = pd.Series(1.0, index=series.index)
        for row in intervals:
            if row["sku"] != sku:
                continue
            start, end = pd.Timestamp(row["start_date"]), pd.Timestamp(row["end_date"])
            availability.loc[(availability.index >= start) & (availability.index <= end)] = 0.0
        history = [{
            "date": str(stamp.date()), "observed_quantity": float(value),
            "availability_fraction": float(availability.loc[stamp]),
            "complete": stamp < pd.Timestamp(origin),
        } for stamp, value in series.items()]
        request = {"schema_version": "forecast-input-v2", "as_of": str(pd.Timestamp(origin).date()),
                   "horizon_days": 28, "items": [{
                       "sku": sku, "unit": "шт", "warehouse_id": "almaty",
                       "supplier_id": "systeme-electric", "category_raw": None,
                       "launch_date": str(series.index[0].date()), "history": history,
                       "events": events.get(sku, []), "known_promotions": [], "analogue_history": [],
                   }]}
        try:
            result, audit = forecast_with_audit(request)
        except ValueError as exc:
            raise DomainError("MISSING_CRITICAL_DATA", f"Контекст регулярного спроса для {sku} противоречив: {exc}.") from exc
        daily[sku], audits[sku] = result[sku], audit[sku]
    return daily, audits


def _history(panel, sku, origin):
    end = pd.Timestamp(origin)
    values = []
    for offset in range(5, -1, -1):
        period_end = end - pd.Timedelta(days=offset * 28)
        period_start = period_end - pd.Timedelta(days=27)
        observed = float(panel.daily.loc[sku, period_start:period_end].sum())
        values.append({"period_start": str(period_start.date()), "period_end": str(period_end.date()),
                       "observed_sales": observed, "regular_sales": None,
                       "estimated_lost_demand": None, "source_kind": "observed"})
    return values


def _policy_maps(request):
    policies = {}
    for row in request["category_policies"]:
        if row["category_raw"] in policies:
            raise DomainError("INVALID_PARAMETERS", "Политика категории передана дважды.")
        minimum = row["minimum_target_quantile"]
        if minimum is not None and row["target_quantile"] < minimum:
            raise DomainError("INVALID_PARAMETERS", "Целевой квантиль категории не может быть ниже её обязательного минимума.")
        policies[row["category_raw"]] = row
    economics = {}
    for row in request["economic_profiles"]:
        if row["sku"] in economics:
            raise DomainError("INVALID_PARAMETERS", "Экономический профиль SKU передан дважды.")
        economics[row["sku"]] = row
    growth = {}
    for row in request["growth_adjustments"]:
        if row["valid_to"] < row["valid_from"]:
            raise DomainError("INVALID_PARAMETERS", "Окончание периода прироста не может предшествовать началу.")
        for previous in growth.get(row["sku"], []):
            if row["valid_from"] <= previous["valid_to"] and previous["valid_from"] <= row["valid_to"]:
                raise DomainError("INVALID_PARAMETERS", "Повторяющиеся или пересекающиеся периоды прироста одного SKU не допускаются; передайте одну согласованную поправку.")
        growth.setdefault(row["sku"], []).append(row)
    return policies, economics, growth


class TirekCalculationPipeline:
    """Use frozen forecast weights and the production decision core."""

    def calculate(self, dataset: dict, request: dict, calculation_id: str) -> dict:
        # Validate contradictions for every supplier adapter before dispatch.
        policies, economics, growth = _policy_maps(request)
        if dataset.get('supplier_ids') == ['iek']:
            from .iek_adapter import IEKForecastPipeline
            return IEKForecastPipeline().calculate(dataset, request, calculation_id)
        data_dir = Path(os.getenv("DATA_DIR", str(ROOT / "data")))
        directory = data_dir / "uploads" / dataset["dataset_id"]
        metadata = json.loads((directory / "normalized" / "metadata.json").read_text(encoding="utf-8"))
        current_profiles = json.loads((directory / "normalized" / "current.json").read_text(encoding="utf-8"))
        context = json.loads((directory / "normalized" / "context.json").read_text(encoding="utf-8"))
        panels, _ = load_panels(directory / metadata["model_input_root"], ("SE",))
        panel = next(value for value in panels if value.name == "SE__pieces")
        origin = pd.Timestamp(request["as_of_date"])
        if origin > panel.daily.columns[-1]:
            raise DomainError("MISSING_CRITICAL_DATA", "История продаж не доходит до даты расчёта.")
        frame, point_forecasts, model_entry, selection_hash = _load_forecast(panel, origin)
        prices, material_rows = _context_maps(context, request["as_of_date"], "almaty")
        transaction_path = None
        if context.get("client_labels"):
            transaction_dir = directory / metadata["model_input_root"] / "Systeme electric"
            transaction_paths = sorted(transaction_dir.glob(PATH_PATTERNS["sales_transactions"]))
            if len(transaction_paths) != 1:
                raise DomainError("MISSING_CRITICAL_DATA", "Не найден единственный источник строк продаж для контекста.")
            transaction_path = transaction_paths[0]
        context_daily, context_audits = _context_forecasts(
            panel, origin, context, transaction_path)
        meta = {
            "calculation_id": calculation_id, "dataset_id": dataset["dataset_id"],
            "dataset_version": dataset["dataset_version"], "data_as_of": dataset["data_as_of"],
            "as_of_date": request["as_of_date"], "mode": request["mode"], "revision": 1,
            "policy_version": hashlib_sha256(json.dumps({k: request[k] for k in (
                "category_policies", "economic_profiles", "growth_adjustments")}, sort_keys=True).encode()).hexdigest()[:12],
            "created_at": now(), "horizon_days": request["horizon_days"], "issues": deepcopy(dataset["issues"]),
        }
        details, approval_constraints = {}, {}
        forecast_by_sku = {str(sku): float(value) for sku, value in zip(frame.sku, point_forecasts)}
        forecast_by_sku.update({sku: float(sum(values)) for sku, values in context_daily.items()})
        # Stock/material-only products must remain reviewable. No zero forecast
        # or fabricated daily history is substituted for an unavailable model.
        skus = list(forecast_by_sku) + sorted((set(current_profiles) | set(material_rows)) - set(forecast_by_sku))
        for sku in skus:
            forecast_total = forecast_by_sku.get(sku)
            context_audit = context_audits.get(sku)
            profile = current_profiles.get(sku)
            category = profile.get("category") if profile else None
            if request["category_codes"] and category not in request["category_codes"]:
                continue
            if request["warehouse_ids"] and not set(request["warehouse_ids"]) & {"almaty", "Алматы"}:
                continue
            material_units = {value["unit"] for value in material_rows.get(sku, [])}
            unit = ("шт" if forecast_total is not None else
                    (profile or {}).get("unit") or (next(iter(material_units)) if len(material_units) == 1 else "unknown"))
            inventory = None
            if profile and profile.get("inventory_as_of") and profile["inventory_as_of"] <= request["as_of_date"]:
                inventory = InventorySnapshot(date.fromisoformat(profile["inventory_as_of"]),
                                              profile.get("on_hand"), profile.get("reserved"),
                                              profile.get("free_stock"), profile.get("source", "snapshot"))
            inbound = []
            inbound_missing = []
            if profile is None or profile.get("inbound_quantity") is None:
                inbound_missing.append("inbound_quantity")
            elif profile["inbound_quantity"] != 0 and not profile.get("inbound_eta"):
                inbound_missing.append("inbound_eta")
            elif profile["inbound_quantity"] != 0:
                inbound.append(Inbound(f"inbound-{sku}", profile["inbound_quantity"],
                                       date.fromisoformat(profile["inbound_eta"]), unit, profile["source"]))
            constraint = None
            if profile:
                constraint = SupplierConstraint(unit, profile.get("min_order_qty"), profile.get("order_multiple"), 1,
                                                profile.get("source", "supplier export"))
            econ_row = economics.get(sku)
            price = prices.get(sku)
            econ = EconomicProfile(econ_row["underage_cost"], econ_row["overage_cost"],
                                   econ_row["horizon_days"], econ_row["currency"], econ_row["rationale"]) if econ_row else None
            policy_row = policies.get(category)
            policy = None
            if policy_row:
                # With economics, target_quantile is only a fallback. Only an
                # explicit minimum_target_quantile may constrain the economic choice.
                level = policy_row["minimum_target_quantile"] if econ else policy_row["target_quantile"]
                if level is not None:
                    policy = ServicePolicy(level, policy_row["rationale"], meta["policy_version"])
            adjustments = [GrowthAdjustment(f"growth-{sku}-{index}", value["rate"],
                                             date.fromisoformat(value["valid_from"]), date.fromisoformat(value["valid_to"]),
                                             value["rationale"]) for index, value in enumerate(growth.get(sku, []))]
            materials = []
            for value in material_rows.get(sku, []):
                uncovered = float(Decimal(str(value["quantity"])) - Decimal(str(value["already_accounted_quantity"])))
                materials.append(MaterialRequirement(value["requirement_id"], uncovered,
                                                     date.fromisoformat(value["needed_at"]), value["unit"],
                                                     value["reference"], hard=True))
            unit_cost = (econ_row.get("unit_cost") if econ_row and econ_row.get("unit_cost") is not None
                         else price.get("unit_cost") if price else None)
            mapped_policies = {category: policy} if category is not None and policy is not None else {}
            decision_input = RecommendationInput(
                sku=sku, warehouse_id="almaty", supplier_id="systeme-electric", unit=unit,
                as_of=date.fromisoformat(request["as_of_date"]), forecast_as_of=date.fromisoformat(request["as_of_date"]),
                scenarios=None if forecast_total is None else
                          [context_daily[sku] if sku in context_daily else _daily_path(panel, sku, origin, forecast_total)],
                forecast_id=(f"{REGULAR_VERSION}:{context_audit['selected_method']}" if context_audit else
                             f"forecast-v2-{selection_hash[:12]}"),
                unit_quantum=1 if unit == "шт" else None,
                inventory=inventory, lead_time_days=request["lead_time_days"],
                review_period_days=request["review_period_days"], constraints=constraint,
                economics=econ, service_policy=None,
                category_id=category if mapped_policies else None, category_policies=mapped_policies,
                inbound=inbound, materials=materials, growth_adjustments=adjustments,
                unit_cost=unit_cost, price_currency="KZT" if unit_cost is not None else None,
                data_version=dataset["dataset_version"], max_snapshot_age_days=0, max_forecast_age_days=0,
            )
            try:
                result = recommend(decision_input)
            except ValueError as exc:
                result = Recommendation(sku, "almaty", "systeme-electric", unit,
                                        date.fromisoformat(request["as_of_date"]), "needs_data", None, None,
                                        f"Расчёт отклонён из-за качества входа: {exc}", "unknown", (),
                                        ("valid_source_values",), unit_cost, "KZT" if unit_cost is not None else None,
                                        {"free_stock": profile.get("free_stock") if profile else None,
                                         "inbound_in_horizon": None, "dated_material_demand": None,
                                         "unit_quantum": 1 if unit == "шт" else None,
                                         "min_order_qty_inventory_units": profile.get("min_order_qty") if profile else None,
                                         "order_multiple_inventory_units": profile.get("order_multiple") if profile else None,
                                        "service_floor_quantity": 0, "hard_material_floor_quantity": 0})
            if inbound_missing:
                # Keep the no-inbound path inspectable as an explicitly
                # provisional upper bound. It cannot become final until a
                # manager confirms that no open inbound exists or reimports a
                # dated inbound quantity.
                provisional_without_inbound = (result.quantity if result.quantity is not None
                                               else result.provisional_quantity)
                result = replace(
                    result, status="needs_data", quantity=None,
                    provisional_quantity=provisional_without_inbound,
                    required_fields=tuple(dict.fromkeys((*result.required_fields, *inbound_missing))),
                    reason="Неизвестны данные о поступлениях: " + ", ".join(inbound_missing)
                           + ". Показана предварительная верхняя оценка при нулевом inbound; "
                             "она заблокирована до явного подтверждения или повторного расчёта.",
                    diagnostics={**result.diagnostics, "inbound_in_horizon": None},
                )
            if forecast_total is None:
                result = replace(result, status="needs_data", quantity=None, provisional_quantity=None,
                                 reason="Нет подтверждённого прогноза для SKU: нужна история или явный прогноз аналога. "
                                        "Материальная потребность и складской профиль сохранены; спрос не заменён нулём.",
                                 required_fields=tuple(dict.fromkeys((*result.required_fields, "forecast_history"))))
            if profile and not profile.get("warehouse_scope_verified", False):
                result = replace(
                    result, status="needs_data" if result.status == "needs_data" else "needs_review", quantity=None,
                    provisional_quantity=result.quantity if result.quantity is not None else result.provisional_quantity,
                    warnings=tuple(dict.fromkeys((*result.warnings, "warehouse_scope_unverified"))),
                    required_fields=tuple(dict.fromkeys((*result.required_fields, "warehouse_scope_confirmation"))),
                )
            if forecast_total is not None:
                # Frozen v2 artifacts contain aggregate errors, not calibrated
                # out-of-sample joint paths. Do not manufacture uncertainty from
                # WAPE or reuse residuals belonging to the older v1 model.
                result = replace(result, status="needs_review" if result.status == "ready" else result.status,
                                 warnings=tuple(dict.fromkeys((*result.warnings, "uncalibrated_single_forecast_path"))))
            if context_audit:
                context_warnings = ["context_adjusted_regular_forecast"]
                if context_audit["zero_exposure_days"]:
                    context_warnings.append("stockout_correction_applied")
                if context_audit["excluded_client_windows"]:
                    context_warnings.append("client_oneoff_excluded")
                result = replace(
                    result, status="needs_review" if result.status == "ready" else result.status,
                    warnings=tuple(dict.fromkeys((*result.warnings, *context_warnings))),
                )
            # Known source quantities remain inspectable even when forecasting
            # failed. They are not a computed demand forecast or approval floor.
            material_quantities = [value.quantity for value in materials
                                   if decision_input.as_of < value.due_at <= decision_input.as_of + timedelta(days=request["horizon_days"])]
            material_uncovered = (float(sum((Decimal(str(value)) for value in material_quantities), Decimal(0)))
                                  if all(np.isfinite(value) and value >= 0 for value in material_quantities) else None)
            item_id = "se-" + hashlib_sha256(sku.encode()).hexdigest()[:12]
            shown_quantity = result.quantity if result.quantity is not None else result.provisional_quantity
            decision_status = "no_order" if result.status == "ready" and shown_quantity == 0 else result.status
            issue_rows = [_issue("REQUIRED_" + field.upper().replace(":", "_"), "error",
                                 "Нужно уточнить: " + field, [sku], profile.get("source") if profile else None)
                          for field in result.required_fields]
            issue_rows += [_issue(value.upper(), "warning", value.replace("_", " "), [sku], "decision-core-1")
                           for value in result.warnings]
            if result.provisional_quantity is not None:
                issue_rows.append(_issue("PROVISIONAL_QUANTITY", "warning",
                                         "Показана предварительная потребность; неизвестные условия партии блокируют утверждение.",
                                         [sku], "decision-core-1"))
            model_id = (f"{REGULAR_VERSION}:{context_audit['selected_method']}" if context_audit else
                        f"forecast-v2-{selection_hash[:12]}")
            forecast = None if forecast_total is None else {
                "horizon_days": 28, "period_start": str((origin + pd.Timedelta(days=1)).date()),
                "period_end": str((origin + pd.Timedelta(days=28)).date()), "method": "ml",
                "model_id": model_id, "mean": forecast_total,
                "p10": None, "p50": None, "p90": None,
                "target_quantile": policy.service_level if policy else None, "target_stock": None,
                "calibration_status": "insufficient_data",
                "note": (("Регулярный прогноз пересчитан по доступным дням и обезличенным клиентским окнам. "
                          if context_audit else "Точечный прогноз наблюдаемых продаж; используется одна дневная траектория по прошлым дням недели. ")
                         + "Артефакты не содержат проверенных совместных сценариев ошибок: вероятностный экономический выбор "
                           "и гарантии service level не подтверждены. Квантили не калиброваны."),
            }
            selected_method = model_entry["selected"]
            if forecast is not None:
                forecast["method"] = ("baseline" if context_audit else
                                      "ml" if (selected_method in MODEL_SPECS or selected_method.startswith(("mix|", "blend:"))) else "baseline")
            order_cost = None if shown_quantity is None or unit_cost is None else float(Decimal(str(shown_quantity)) * Decimal(str(unit_cost)))
            # The adapter has not called a provider. The pipeline owns the
            # optional batch review and decides whether configuration is available.
            ai = {"status": "not_requested", "verdict": None, "reasons": [],
                  "evidence_ids": [], "rule_ids": [],
                  "suggested_action": None, "provider_model": None}
            evidence = ([
                {"id": f"{item_id}-forecast", "source_kind": "observed", "reference": model_id,
                 "label": "Прогноз на 28 дней", "value": forecast_total, "unit": "шт"},
            ] if forecast_total is not None else []) + [
                {"id": f"{item_id}-stock", "source_kind": "observed", "reference": profile.get("source") if profile else "missing-current-profile",
                 "label": "Свободный остаток", "value": result.diagnostics.get("free_stock"), "unit": unit if unit != "unknown" else None},
            ]
            item = {
                "item_id": item_id, "supplier_id": "systeme-electric", "supplier_name": "Systeme Electric",
                "sku": sku, "supplier_article": profile.get("supplier_article") if profile else None,
                "name": profile.get("name", sku) if profile else sku, "unit": unit, "warehouse_id": "almaty",
                "category_raw": category, "policy_basis": "economic" if econ else "service_policy" if policy else None,
                "decision_status": decision_status, "urgency": "critical" if result.urgency == "expedite" else result.urgency,
                "free_stock": result.diagnostics.get("free_stock"),
                "inbound_within_horizon": result.diagnostics.get("inbound_in_horizon"),
                "material_requirement_uncovered": material_uncovered,
                "forecast": forecast, "min_order_qty": result.diagnostics.get("min_order_qty_inventory_units"),
                "order_multiple": result.diagnostics.get("order_multiple_inventory_units"),
                "recommended_quantity": shown_quantity, "final_quantity": result.quantity,
                "override_reason": None, "unit_cost_kzt": unit_cost, "order_cost_kzt": order_cost if result.quantity is not None else None,
                "marginal_value": None, "economics_source": econ_row["source_kind"] if econ_row else None,
                "reason": result.reason, "issues": issue_rows, "evidence": evidence, "ai": ai,
            }
            applied = (["POLICY-01"] if econ else ["POLICY-02"] if policy else ["DATA-01"])
            if context_audit:
                if context_audit["zero_exposure_days"]:
                    applied.append("DEMAND-01")
                if context_audit["excluded_client_windows"]:
                    applied.append("DEMAND-02")
            details[item_id] = {"meta": meta, "item": item, "history": _history(panel, sku, origin) if forecast_total is not None else [],
                                "inbound": [{"order_id": value.event_id, "quantity": value.quantity,
                                             "expected_at": str(value.expected_at), "source_reference": value.source}
                                            for value in inbound],
                                "economic_profile": deepcopy(econ_row), "applied_rule_ids": applied}
            approval_constraints[item_id] = {
                "unit_quantum": result.diagnostics.get("unit_quantum"),
                "min_order_qty": result.diagnostics.get("min_order_qty_inventory_units"),
                "order_multiple": result.diagnostics.get("order_multiple_inventory_units"),
                "minimum_safe_quantity": max(result.diagnostics.get("service_floor_quantity", 0),
                                             result.diagnostics.get("hard_material_floor_quantity", 0)),
                "constraints_complete": result.diagnostics.get("min_order_qty_inventory_units") is not None
                                        and result.diagnostics.get("order_multiple_inventory_units") is not None
                                        and bool(profile and profile.get("warehouse_scope_verified")),
            }
        items = [value["item"] for value in details.values()]
        return {"response": {"meta": meta, "summary": summary(items), "items": items},
                "details": details, "request": request, "approval_constraints": approval_constraints}
