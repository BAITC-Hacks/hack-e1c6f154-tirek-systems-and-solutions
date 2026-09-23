"""Observed IEK sales -> frozen forecast. Missing current stock is never invented."""
from copy import deepcopy
from hashlib import sha256
import json
import os
from pathlib import Path
import shutil

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from model.forecast_v2.data import code, load_panels
from .contracts import DomainError, ROOT
from .demo import now, summary
from .ml_adapter import _history, _issue, _load_forecast, _write_json


def normalize_dataset(directory, report, context):
    directory = Path(directory)
    by_role = {source['role']: source for source in report['sources']}
    required = {'sales_transactions', 'sales_monthly'}
    if not required.issubset(by_role):
        report['issues'].append(_issue('NORMALIZATION_BLOCKED', 'error',
                                      'Для прогноза IEK нужны динамика и помесячные продажи.'))
        return report
    folder = directory / 'model_inputs' / 'IEK'
    folder.mkdir(parents=True, exist_ok=True)
    for role, source in by_role.items():
        filename = Path(source['filename']).name
        path = directory / f'{role}.xlsx'
        if sha256(path.read_bytes()).hexdigest() != source['sha256']:
            raise DomainError('INVALID_FILE', 'Контрольная сумма исходного файла изменилась.')
        shutil.copyfile(path, folder / filename)
    try:
        panels, audits = load_panels(folder.parent, ('IEK',))
    except (ValueError, KeyError, OSError) as exc:
        raise DomainError('INVALID_FILE', 'Формат выгрузки IEK не соответствует динамике и помесячным продажам.') from exc
    names = {}
    workbook = load_workbook(folder / by_role['sales_transactions']['filename'], read_only=True, data_only=True)
    try:
        for row in workbook.worksheets[0].iter_rows(min_row=2, values_only=True):
            if len(row) >= 8 and row[3] is not None:
                names[code(row[3])] = str(row[4] or row[3])
    finally:
        workbook.close()
    normalized = directory / 'normalized'
    normalized.mkdir(exist_ok=True)
    _write_json(normalized / 'names.json', names)
    _write_json(normalized / 'audit.json', audits['IEK'])
    audit = audits['IEK']
    as_of = audit['transactions']['data_as_of']
    for source in report['sources']:
        if source['role'] == 'sales_transactions':
            source.update(rows_used=audit['transactions']['used_positive_invoice_rows'], data_as_of=as_of)
        elif source['role'] == 'sales_monthly':
            source['rows_used'] = audit['monthly']['sku_count']
    report.update(data_as_of=as_of, sku_count=sum(len(panel.daily) for panel in panels), calculation_allowed=True)
    report['issues'] = [issue for issue in report['issues']
                        if issue['code'] not in {'NORMALIZATION_REQUIRED', 'SOURCE_DATE_UNKNOWN'}]
    report['issues'].extend([
        _issue('FORECAST_ONLY', 'warning', 'Прогноз IEK рассчитан по наблюдаемым продажам. Актуального свободного остатка нет: количество закупки останется неизвестным.'),
        _issue('AUXILIARY_SOURCES_NOT_APPLIED', 'warning', 'Путь, месячные остатки, сезонность и MOQ сохранены. В этом адаптере они не подменяют актуальный остаток и не участвуют в расчёте заказа.'),
        _issue('NO_DAILY_STOCKOUT', 'warning', 'Продажи не скорректированы на stockout и клиентские разовые заказы.'),
        _issue('LAST_DAY_UNVERIFIED', 'warning', 'Полнота последнего дня выгрузки не подтверждена. Он используется как доступная часть дня.'),
    ])
    if context:
        report['issues'].append(_issue('CONTEXT_NOT_APPLIED', 'warning', 'Дополнительный контекст сохранён; адаптер прогноза IEK пока не использует его для заказа.'))
        if any(row.get('source_kind') == 'synthetic' for key in ('stockout_intervals', 'price_observations', 'material_requirements') for row in context.get(key, [])):
            report['issues'].append(_issue('SYNTHETIC_CONTEXT_BLOCKS_OPERATIONAL', 'warning', 'Синтетический контекст допускает только сценарный режим.'))
    return report


class IEKForecastPipeline:
    def calculate(self, dataset, request, calculation_id):
        if request['category_codes'] or any(request[key] for key in ('category_policies', 'economic_profiles', 'growth_adjustments')):
            raise DomainError('INVALID_PARAMETERS', 'IEK сейчас поддерживает прогноз продаж без категорий, экономических политик и прироста. Для закупки сначала нужен актуальный складской снимок.')
        if set(request['warehouse_ids']) - {'almaty'}:
            raise DomainError('INVALID_PARAMETERS', 'В выгрузке IEK поддержан склад Алматы.')
        directory = Path(os.getenv('DATA_DIR', str(ROOT / 'data'))) / 'uploads' / dataset['dataset_id']
        names = json.loads((directory / 'normalized' / 'names.json').read_text(encoding='utf-8'))
        panels, _ = load_panels(directory / 'model_inputs', ('IEK',))
        origin = pd.Timestamp(request['as_of_date'])
        meta = {'calculation_id': calculation_id, 'dataset_id': dataset['dataset_id'],
                'dataset_version': dataset['dataset_version'], 'data_as_of': dataset['data_as_of'],
                'as_of_date': request['as_of_date'], 'mode': request['mode'], 'revision': 1,
                'policy_version': 'iek-forecast-only-v1', 'created_at': now(), 'horizon_days': 28,
                'issues': deepcopy(dataset['issues'])}
        details = {}
        for panel in panels:
            if origin > panel.daily.columns[-1]:
                raise DomainError('MISSING_CRITICAL_DATA', 'История IEK не доходит до даты расчёта.')
            frame, forecasts, entry, selection_hash = _load_forecast(panel, origin)
            if not np.isfinite(forecasts).all() or (forecasts < 0).any():
                raise DomainError('INVALID_PIPELINE_RESULT', 'Модель вернула некорректный прогноз.', 500)
            for position, sku in enumerate(frame.sku.astype(str)):
                value = float(forecasts[position])
                item_id = 'iek-' + sha256((sku + '|' + panel.unit).encode()).hexdigest()[:12]
                item = {
                    'item_id': item_id, 'supplier_id': 'iek', 'supplier_name': 'IEK', 'sku': sku,
                    'supplier_article': None, 'name': names.get(sku, sku), 'unit': panel.unit,
                    'warehouse_id': 'almaty', 'category_raw': None, 'policy_basis': None,
                    'decision_status': 'needs_data', 'urgency': 'unknown', 'free_stock': None,
                    'inbound_within_horizon': None, 'material_requirement_uncovered': None,
                    'forecast': {'horizon_days': 28, 'period_start': str((origin + pd.Timedelta(days=1)).date()),
                                 'period_end': str((origin + pd.Timedelta(days=28)).date()),
                                 'method': 'baseline' if entry['selected'].startswith(('mean', 'median', 'snaive', 'croston', 'tsb')) else 'ml',
                                 'model_id': f"forecast-v2-{selection_hash[:12]}:{panel.name}", 'mean': value,
                                 'p10': None, 'p50': None, 'p90': None, 'target_quantile': None, 'target_stock': None,
                                 'calibration_status': 'insufficient_data',
                                 'note': f"Прогноз наблюдаемых продаж; метод {entry['selected']}. Интервалы не калиброваны."},
                    'min_order_qty': None, 'order_multiple': None, 'recommended_quantity': None,
                    'final_quantity': None, 'override_reason': None, 'unit_cost_kzt': None,
                    'order_cost_kzt': None, 'marginal_value': None, 'economics_source': None,
                    'reason': 'ML-прогноз готов. Для количества закупки нужен актуальный свободный остаток, условия поставки и политика запаса.',
                    'issues': [_issue('CURRENT_STOCK_REQUIRED', 'error', 'Нет актуального свободного остатка IEK.', [sku])],
                    'evidence': [{'id': item_id + '-forecast', 'source_kind': 'observed',
                                  'reference': f"forecast-v2:{selection_hash}:{panel.name}",
                                  'label': 'Прогноз по наблюдаемым продажам на 28 дней', 'value': value, 'unit': panel.unit}],
                    'ai': {'status': 'not_requested', 'verdict': None, 'reasons': [], 'evidence_ids': [],
                           'rule_ids': [], 'suggested_action': None, 'provider_model': None},
                }
                details[item_id] = {'meta': meta, 'item': item, 'history': _history(panel, sku, origin),
                                    'inbound': [], 'economic_profile': None, 'applied_rule_ids': ['DATA-01']}
        items = [detail['item'] for detail in details.values()]
        return {'response': {'meta': meta, 'summary': summary(items), 'items': items},
                'details': details, 'request': request, 'approval_constraints': {}}
