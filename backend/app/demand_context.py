"""Join explicit workbook events/stockouts to the regular-demand forecaster.

The uploaded sales ledger remains immutable. Row IDs include role, sheet and the
one-based Excel row number, so repeated invoices do not collapse into one event.
No customer identity or stockout is inferred from monthly stock snapshots.
"""
from collections import defaultdict
from datetime import date

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from model.forecast_v2.data import code
from model.regular_forecast.predictor import VERSION, clean_history, forecast_with_audit

from .contracts import DomainError


def normalize_events(path, context):
    """Resolve supplied client labels to positive invoice rows; reject bad joins."""
    labels = {}
    for label in (context or {}).get("client_labels", []):
        event_id = label["source_event_id"]
        if event_id in labels:
            raise ValueError("Duplicate client label: " + event_id)
        labels[event_id] = label["pseudonymous_client_id"]
    if not labels:
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    events, matched = [], set()
    try:
        sheet = workbook.worksheets[0]
        for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            event_id = f"sales_transactions:{sheet.title}:{number}"
            document = str(row[1]).strip() if row[1] is not None else None
            aliases = {event_id}
            if document:
                aliases.update((document, f"{document}:{number}"))
            keys = aliases & labels.keys()
            if not keys:
                continue
            if len({labels[key] for key in keys}) != 1:
                raise ValueError("Conflicting client labels for the same invoice row: " + event_id)
            stamp = pd.to_datetime(row[0], format="mixed", dayfirst=True, errors="coerce")
            quantity = row[7]
            if (pd.isna(stamp) or stamp < pd.Timestamp("2025-01-01") or row[3] is None
                    or not str(row[2]).startswith("Расходная накладная")
                    or isinstance(quantity, bool) or not isinstance(quantity, (int, float))
                    or not np.isfinite(quantity) or quantity <= 0):
                raise ValueError("Client label does not reference a usable positive invoice: " + event_id)
            matched.update(keys)
            events.append({"event_id": event_id, "date": str(stamp.date()),
                           "client_id": labels[next(iter(keys))], "quantity": float(quantity),
                           "sku": code(row[3]), "unit": str(row[5]),
                           "warehouse_id": "almaty" if row[6] in ("Алматы", "almaty") else str(row[6])})
    finally:
        workbook.close()
    unmatched = set(labels) - matched
    if unmatched:
        raise ValueError("Unknown client source_event_id; use sales_transactions:<sheet>:<Excel row>, document or document:row: "
                         + ", ".join(sorted(unmatched)[:3]))
    return events


def validate_intervals(context, panels):
    """Stockouts must describe observed history and cannot contradict a sale."""
    by_sku = {str(sku): panel for panel in panels for sku in panel.daily.index}
    for interval in (context or {}).get("stockout_intervals", []):
        panel = by_sku.get(interval["sku"])
        if interval["warehouse_id"] not in ("almaty", "Алматы"):
            # Other warehouses may carry materials/context but cannot correct this panel.
            continue
        if panel is None:
            raise ValueError("Stockout refers to an unknown sales SKU: " + interval["sku"])
        start, end = pd.Timestamp(interval["start_date"]), pd.Timestamp(interval["end_date"])
        if start > end or start < panel.daily.columns[0] or end > panel.daily.columns[-1]:
            raise ValueError("Stockout interval is reversed or outside observed history: " + interval["sku"])
        if panel.daily.loc[interval["sku"], start:end].sum() > 0:
            raise ValueError("Positive sales contradict a full-day stockout: " + interval["sku"])


def regular_forecasts(panel, origin, context, events, categories=None):
    """Return corrected forecasts only where explicit supplemental context exists.

    Outside supplied stockout days, full availability is an assumption, recorded
    in the explanation. The source's unverified final day is excluded from GLM
    fitting. Unmatched SKU/warehouse events cannot silently affect another item.
    """
    context = context or {}
    cutoff = str(pd.Timestamp(origin).date())
    by_sku, intervals = defaultdict(list), defaultdict(list)
    for event in events:
        if event["unit"] == panel.unit and event["warehouse_id"] == "almaty" and event["date"] <= cutoff:
            by_sku[event["sku"]].append({key: event[key] for key in ("event_id", "date", "client_id", "quantity")})
    for interval in context.get("stockout_intervals", []):
        if interval["warehouse_id"] in ("almaty", "Алматы") and interval["start_date"] <= cutoff:
            intervals[interval["sku"]].append(interval)
    corrected = {}
    for sku in sorted((set(by_sku) | set(intervals)) & set(panel.daily.index)):
        observed = panel.daily.loc[sku, :pd.Timestamp(origin)]
        positive = observed[observed > 0]
        if positive.empty:
            continue
        # Calendar padding before the first recorded sale is not product history.
        first = min([positive.index[0], *[pd.Timestamp(i["start_date"]) for i in intervals[sku]]])
        observed = observed.loc[first:]
        history = []
        for stamp, value in observed.items():
            iso = str(stamp.date())
            unavailable = any(i["start_date"] <= iso <= i["end_date"] for i in intervals[sku])
            history.append({"date": iso, "observed_quantity": float(value),
                            "availability_fraction": 0.0 if unavailable else 1.0,
                            "complete": stamp != panel.daily.columns[-1]})
        item = {"sku": sku, "unit": panel.unit, "warehouse_id": "almaty", "supplier_id": panel.supplier,
                "category_raw": (categories or {}).get(sku), "launch_date": str(first.date()),
                "history": history, "events": by_sku[sku], "known_promotions": [], "analogue_history": []}
        request = {"schema_version": "forecast-input-v2", "as_of": cutoff, "horizon_days": 28, "items": [item]}
        try:
            predictions, audits = forecast_with_audit(request)
            stamps, cleaned, _, _ = clean_history(item, cutoff)
        except ValueError as exc:
            raise DomainError("MISSING_CRITICAL_DATA", f"Регулярный спрос {sku}: {exc}", 422) from exc
        audit = audits[sku]
        excluded = sum(window["quantity"] for window in audit["excluded_client_windows"])
        stockout_days = sum(1 for row in history if row["availability_fraction"] == 0)
        corrected[sku] = {"daily": predictions[sku], "mean": float(sum(predictions[sku])),
                          "audit": audit, "regular_history": dict(zip(stamps, cleaned.tolist())),
                          "excluded_quantity": excluded, "stockout_days": stockout_days,
                          "model_id": f"{VERSION}:{audit['selected_method']}",
                          "note": (f"Регулярный спрос: {audit['selected_method']}; учтено дней отсутствия: {stockout_days}, "
                                   f"исключено разового клиентского объёма: {excluded:g}. "
                                   "Вне переданных интервалов принята полная доступность; последний непроверенный день "
                                   "не участвует в обучении. Клиентская очистка — объяснимая эвристика; квантили не калиброваны.")}
    return corrected


def context_issues(context):
    """Dataset-level capabilities; SKU-level quantities are exposed separately."""
    context = context or {}
    return [
        ("STOCKOUT_CONTEXT_APPLIED" if context.get("stockout_intervals") else "NO_DAILY_STOCKOUT",
         "Переданные интервалы отсутствия используются в регулярном прогнозе; вне них принята полная доступность."
         if context.get("stockout_intervals") else
         "Точных интервалов отсутствия товара нет; скрытый спрос не восстановлен."),
        ("CLIENT_CONTEXT_APPLIED" if context.get("client_labels") else "NO_CLIENT_LABELS",
         "Клиентские метки связаны со строками XLSX; крупные разовые покупки исключаются объяснимой эвристикой."
         if context.get("client_labels") else
         "Обезличенные клиентские метки отсутствуют; разовые клиентские заказы не классифицированы."),
    ]
