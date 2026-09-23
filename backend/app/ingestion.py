"""Safe workbook inspection/storage; partner-specific normalization is intentionally separate."""
from datetime import date
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
from .contracts import DomainError, validate
from .demo import ROLES

MAX_BYTES = 30 * 1024 * 1024


def detect_role(filename):
    name = filename.lower()
    for role in ROLES:
        if role in name:
            return role
    for fragment, role in [('динамика', 'sales_transactions'), ('продаж', 'sales_monthly'),
                           ('остат', 'stock_monthly'), ('сезон', 'seasonality'), ('moq', 'moq'), ('пути', 'current_stock_inbound'), ('путь', 'current_stock_inbound')]:
        if fragment in name:
            return role
    raise DomainError('INVALID_FILE', f'Не определена роль файла «{filename}». Используйте названия ролей из шаблона.', 400)


def inspect_upload(files, supplier_id, context):
    if supplier_id not in ('systeme-electric', 'iek'):
        raise DomainError('INVALID_PARAMETERS', 'Поддержаны Systeme Electric и IEK.')
    if not 1 <= len(files) <= 6:
        raise DomainError('INVALID_FILE', 'Загрузите от одного до шести файлов XLSX.', 400)
    if sum(len(content) for _, content in files) > MAX_BYTES:
        raise DomainError('UPLOAD_TOO_LARGE', 'Общий размер файлов превышает 30 МБ.', 413)
    if context is not None:
        validate('AdditionalContext', context)
    sources, roles = [], set()
    for filename, content in files:
        if Path(filename).suffix.lower() != '.xlsx':
            raise DomainError('INVALID_FILE', 'Разрешены только файлы .xlsx.', 400)
        role = detect_role(filename)
        if role in roles:
            raise DomainError('INVALID_FILE', f'Роль «{ROLES[role]}» указана дважды.', 400)
        roles.add(role)
        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 120 * 1024 * 1024:
                    raise DomainError('UPLOAD_TOO_LARGE', 'Распакованный XLSX превышает 120 МБ.', 413)
            book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
            rows = 0
            for sheet in book:
                if (sheet.max_row or 0) > 250000 or (sheet.max_column or 0) > 256:
                    raise DomainError('INVALID_FILE', 'Лист превышает лимит 250 000 строк / 256 столбцов.', 400)
                rows += sum(1 for row in sheet.iter_rows(values_only=True) if any(cell is not None for cell in row))
            book.close()
        except DomainError:
            raise
        except Exception as exc:
            raise DomainError('INVALID_FILE', f'Не удалось прочитать XLSX «{filename}». Проверьте файл.', 400) from exc
        sources.append({'role': role, 'filename': Path(filename).name, 'sha256': sha256(content).hexdigest(),
                        'rows_read': rows, 'rows_used': 0, 'data_as_of': None})
    # The snapshot date can be encoded in the original filename. Identical
    # bytes with different dated names must not overwrite the same dataset.
    fingerprint = sha256(json.dumps({'supplier': supplier_id, 'sources': sorted([(s['role'], s['filename'], s['sha256']) for s in sources]),
                                      'context': context}, sort_keys=True).encode()).hexdigest()
    issues = [{'code': 'NORMALIZATION_REQUIRED', 'severity': 'warning',
               'message': 'Файлы проверены и сохранены. Даты, коды и строки ещё не нормализованы; прогноз не запускался.',
               'affected_skus': [], 'source_reference': None}]
    for missing in ROLES.keys() - roles:
        issues.append({'code': 'MISSING_SOURCE', 'severity': 'error', 'message': f'Нет источника: {ROLES[missing]}.',
                       'affected_skus': [], 'source_reference': missing})
    report = {'dataset_id': 'ds-' + fingerprint[:20], 'dataset_version': fingerprint[:12],
              'data_as_of': str(date.today()), 'timezone': 'Asia/Almaty', 'source_kind': 'observed',
              'supplier_ids': [supplier_id], 'sku_count': 0, 'sources': sources, 'issues': issues, 'calculation_allowed': False}
    # Required contract date is an upload date until the adapter can establish source freshness.
    issues.append({'code': 'SOURCE_DATE_UNKNOWN', 'severity': 'warning',
                   'message': 'Дата источников не установлена. Дата набора означает дату загрузки.',
                   'affected_skus': [], 'source_reference': None})
    return report
