"""Optional explicit current-stock template for IEK procurement.

Legacy inbound-only exports do not satisfy this schema. They remain available
for provenance but never manufacture a current stock balance or supplier MOQ.
"""
from datetime import date, datetime
from math import isfinite
from pathlib import Path

from openpyxl import load_workbook


HEADERS = ("Код", "Номенклатура", "Ед.", "Склад", "Дата снимка", "Остаток",
           "Зарезервировано", "Свободный остаток", "Минимальная партия", "Кратность",
           "Шаг единицы", "Категория", "Цена закупки KZT", "В пути", "Дата поступления")


def _date(value, field):
    if isinstance(value, datetime):
        return str(value.date())
    if isinstance(value, date):
        return str(value)
    try:
        return str(date.fromisoformat(str(value)))
    except ValueError as exc:
        raise ValueError(f"{field}: use an Excel date or YYYY-MM-DD") from exc


def _quantity(value, field, strictly_positive=False):
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not isfinite(value):
        raise ValueError(f"{field}: explicit finite numeric value required")
    if value < 0 or (strictly_positive and value == 0):
        raise ValueError(f"{field}: invalid negative or zero quantity")
    return float(value)


def read_current_stock(path):
    """Return None for a legacy non-template workbook; reject malformed templates."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        header = next(sheet.iter_rows(values_only=True), ())
        if not header or header[0] != "Код":
            return None
        if any(header.count(name) != 1 for name in HEADERS):
            raise ValueError("Current stock template must contain each documented header exactly once")
        positions = {name: header.index(name) for name in HEADERS}
        profiles = {}
        for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(value is not None for value in row):
                continue
            values = {name: row[position] for name, position in positions.items()}
            sku = values["Код"]
            if not isinstance(sku, str) or not sku.strip():
                raise ValueError("Current stock SKU must be nonempty text preserving leading zeros")
            sku = sku.strip()
            if sku in profiles:
                raise ValueError("Duplicate current stock SKU: " + sku)
            if values["Склад"] not in ("almaty", "Алматы"):
                raise ValueError("Current stock template supports warehouse almaty / Алматы")
            if values["Ед."] not in ("шт", "м", "упак"):
                raise ValueError("Current stock unit must be шт, м or упак")
            as_of = _date(values["Дата снимка"], "Дата снимка")
            inbound = _quantity(values["В пути"], "В пути")
            eta = _date(values["Дата поступления"], "Дата поступления") if values["Дата поступления"] is not None else None
            if inbound is not None and inbound > 0 and (eta is None or eta < as_of):
                raise ValueError("Positive inbound requires an explicit ETA on or after the stock snapshot")
            # Explicit zero inbound means no receipt; a missing amount stays unknown.
            on_hand = _quantity(values["Остаток"], "Остаток")
            reserved = _quantity(values["Зарезервировано"], "Зарезервировано")
            free = _quantity(values["Свободный остаток"], "Свободный остаток")
            if all(value is not None for value in (on_hand, reserved, free)) and abs(on_hand - reserved - free) > 1e-8:
                raise ValueError("Current stock free balance contradicts on_hand minus reserved")
            profiles[sku] = {
                "sku": sku, "name": str(values["Номенклатура"] or sku), "supplier_id": "iek",
                "supplier_article": None, "unit": values["Ед."], "warehouse_id": "almaty",
                "warehouse_scope_verified": True, "inventory_as_of": as_of,
                "on_hand": on_hand, "reserved": reserved, "free_stock": free,
                "min_order_qty": _quantity(values["Минимальная партия"], "Минимальная партия"),
                "order_multiple": _quantity(values["Кратность"], "Кратность", True),
                "unit_quantum": _quantity(values["Шаг единицы"], "Шаг единицы", True),
                "category": str(values["Категория"]).strip() if values["Категория"] is not None else None,
                "unit_cost": _quantity(values["Цена закупки KZT"], "Цена закупки KZT"),
                "inbound_quantity": inbound, "inbound_eta": eta,
                "source": f"{Path(path).name}:{sheet.title}:{number}",
            }
        if not profiles:
            raise ValueError("Current stock template has no item rows")
        return profiles
    finally:
        book.close()
